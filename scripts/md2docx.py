# -*- coding: utf-8 -*-
"""将本教学设计 Markdown 转换为 docx（python-docx）。"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\HuaweiMoveData\Users\Herry\Documents\kimi\tasks\2026-10-13\12-09-06-fa1b1707\滑坡与泥石流_单元与课时教学设计.md"
OUT = r"D:\HuaweiMoveData\Users\Herry\Documents\kimi\tasks\2026-10-13\12-09-06-fa1b1707\滑坡与泥石流_单元与课时教学设计.docx"

doc = Document()

# ---------- 全局样式 ----------
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.4
normal.paragraph_format.space_after = Pt(4)

def set_east_asia(style_or_run, name):
    rPr = style_or_run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)

for lvl, size, color in [(1, 16, "1F4E79"), (2, 14, "1F4E79"), (3, 12, "2E74B5")]:
    st = doc.styles[f"Heading {lvl}"]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    set_east_asia(st, "微软雅黑")
    st.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
    st.paragraph_format.space_after = Pt(6)

for sec in doc.sections:
    sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6); sec.right_margin = Cm(2.6)

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

def add_runs(par, text):
    """解析 **bold** 与 `code` 行内标记。"""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            _add_code_aware(par, text[pos:m.start()], False)
        r = par.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        _add_code_aware(par, text[pos:], False)

def _add_code_aware(par, text, bold):
    parts = re.split(r"(`[^`]+`)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("`") and p.endswith("`"):
            r = par.add_run(p[1:-1])
            r.font.name = "Consolas"
            set_east_asia(r, "微软雅黑")
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor.from_string("C00000")
        else:
            r = par.add_run(p)
            r.bold = bold

def add_hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "999999")
    pbdr.append(bottom); pPr.append(pbdr)
    p.paragraph_format.space_after = Pt(8)

def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell_text in enumerate(header):
        cell = t.cell(0, j)
        cell.paragraphs[0].text = ""
        add_runs(cell.paragraphs[0], cell_text)
        for r in cell.paragraphs[0].runs:
            r.bold = True
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "D9E2F3")
        cell._tc.get_or_add_tcPr().append(shd)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            txt = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], txt)
            for par in cell.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(10)
    for j in range(len(header)):
        t.cell(0, j).paragraphs[0].runs and None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def flush_table(buf):
    rows = []
    for line in buf:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c or "---") for c in cells):
            continue  # 分隔行
        rows.append(cells)
    if rows:
        add_table(rows)

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

in_code = False
table_buf = []
first_h1 = True

for raw in lines:
    line = raw.rstrip()
    if line.strip().startswith("```"):
        in_code = not in_code
        continue
    if in_code:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.name = "Consolas"; set_east_asia(r, "微软雅黑")
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(0)
        continue
    if line.strip().startswith("|"):
        table_buf.append(line)
        continue
    elif table_buf:
        flush_table(table_buf); table_buf = []

    s = line.strip()
    if not s:
        continue
    if s == "---":
        add_hr(); continue
    m = re.match(r"^(#{1,6})\s+(.*)$", s)
    if m:
        lvl = min(len(m.group(1)), 3)
        h = doc.add_heading("", level=lvl)
        add_runs(h, m.group(2))
        if lvl == 1 and first_h1:
            first_h1 = False
        continue
    if s.startswith("> "):
        p = doc.add_paragraph(style="Intense Quote") if "Intense Quote" in [st.name for st in doc.styles] else doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        add_runs(p, s[2:])
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string("404040")
        continue
    m = re.match(r"^(\d+)\.\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, m.group(2)); continue
    m = re.match(r"^[-*]\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, m.group(1)); continue
    p = doc.add_paragraph()
    add_runs(p, s)

if table_buf:
    flush_table(table_buf)

# 页脚页码
for sec in doc.sections:
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

doc.save(OUT)
print("saved:", OUT)
